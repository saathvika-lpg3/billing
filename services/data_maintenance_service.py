from __future__ import annotations

import json
import shutil
import sqlite3
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from config.app_config import AppConfig


@dataclass(frozen=True)
class FreeSpaceResult:
    backup_path: Path
    export_path: Path
    records_exported: int
    records_deleted: int
    table_counts: dict[str, int]


class DataMaintenanceService:
    PARENT_SPECS: tuple[dict[str, Any], ...] = (
        {"table": "sales", "date": "bill_date", "children": (("sales_items", "sale_id"), ("einvoices", "sale_id"), ("eway_bills", "sale_id"))},
        {"table": "purchases", "date": "bill_date", "children": (("purchase_items", "purchase_id"),)},
        {"table": "sales_returns", "date": "return_date", "children": (("sales_return_items", "return_id"),)},
        {"table": "purchase_returns", "date": "return_date", "children": (("purchase_return_items", "return_id"),)},
        {"table": "dispatch_returns", "date": "return_date", "children": (("dispatch_return_items", "dispatch_return_id"),)},
        {"table": "order_documents", "date": "doc_date", "children": (("order_document_items", "order_document_id"),)},
        {"table": "stock_outs", "date": "challan_date", "children": (("stock_out_items", "stock_out_id"),)},
        {"table": "stock_transfers", "date": "transfer_date", "children": (("stock_transfer_items", "stock_transfer_id"),)},
        {"table": "production_runs", "date": "production_date", "children": (("production_run_items", "production_run_id"),)},
        {"table": "voucher_headers", "date": "voucher_date", "children": (("voucher_lines", "voucher_id"),)},
    )

    DATE_SPECS: tuple[tuple[str, str], ...] = (
        ("bank_reconciliations", "voucher_date"),
        ("business_registers", "record_date"),
        ("ca_export_logs", "from_date"),
        ("closing_periods", "from_date"),
        ("document_archive_logs", "created_at"),
        ("expenses", "expense_date"),
        ("gst_adjustments", "adjustment_date"),
        ("gst_api_logs", "created_at"),
        ("gst_itc_reconciliations", "invoice_date"),
        ("gst_postings", "voucher_date"),
        ("gst_return_filings", "period_from"),
        ("journal_entries", "entry_date"),
        ("ledger_postings", "voucher_date"),
        ("payments", "payment_date"),
        ("print_batches", "created_at"),
        ("print_logs", "created_at"),
        ("receipts", "receipt_date"),
        ("route_settlements", "settlement_date"),
        ("stock_adjustments", "adj_date"),
        ("stock_log", "entry_date"),
        ("stock_postings", "voucher_date"),
        ("stock_valuation_adjustments", "adjustment_date"),
        ("audit_logs", "created_at"),
    )

    def __init__(self, config: AppConfig, db_path: Path | str) -> None:
        self.config = config
        self.db_path = Path(db_path)
        self.archive_root = self.config.project_root / "archives" / "free_space"

    def free_space_by_year_range(self, from_year: int, to_year: int, *, user_name: str = "Developer") -> FreeSpaceResult:
        if from_year > to_year:
            raise ValueError("From year cannot be greater than To year.")
        start_date = f"{from_year:04d}-01-01"
        end_date = f"{to_year:04d}-12-31"
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        self.archive_root.mkdir(parents=True, exist_ok=True)
        backup_path = self.archive_root / f"sqlite_full_backup_before_free_space_{from_year}_{to_year}_{stamp}.db"
        export_path = self.archive_root / f"free_space_export_{from_year}_{to_year}_{stamp}.xlsx"
        shutil.copy2(self.db_path, backup_path)

        with sqlite3.connect(self.db_path) as conn:
            conn.row_factory = sqlite3.Row
            table_rows = self._collect_rows(conn, start_date, end_date)
            self._export_excel(table_rows, export_path, from_year, to_year)
            records_exported = sum(len(rows) for rows in table_rows.values())
            records_deleted = self._delete_rows(conn, table_rows, start_date, end_date)
            self._log(conn, from_year, to_year, backup_path, export_path, records_exported, records_deleted, user_name, table_rows)
            conn.commit()
        with sqlite3.connect(self.db_path) as vacuum_conn:
            vacuum_conn.execute("VACUUM")

        return FreeSpaceResult(
            backup_path=backup_path,
            export_path=export_path,
            records_exported=records_exported,
            records_deleted=records_deleted,
            table_counts={table: len(rows) for table, rows in table_rows.items() if rows},
        )

    def _collect_rows(self, conn: sqlite3.Connection, start_date: str, end_date: str) -> dict[str, list[dict[str, Any]]]:
        existing = self._tables(conn)
        table_rows: dict[str, list[dict[str, Any]]] = {}
        for spec in self.PARENT_SPECS:
            table = spec["table"]
            if table not in existing:
                continue
            parents = self._rows_between(conn, table, spec["date"], start_date, end_date)
            if parents:
                table_rows[table] = parents
            ids = [row["id"] for row in parents if row.get("id") is not None]
            for child_table, fk in spec["children"]:
                if child_table not in existing or not ids:
                    continue
                child_rows = self._rows_by_ids(conn, child_table, fk, ids)
                if child_rows:
                    table_rows[child_table] = child_rows
        for table, date_column in self.DATE_SPECS:
            if table not in existing:
                continue
            rows = self._rows_between(conn, table, date_column, start_date, end_date)
            if rows:
                table_rows[table] = rows
        return table_rows

    def _delete_rows(self, conn: sqlite3.Connection, table_rows: dict[str, list[dict[str, Any]]], start_date: str, end_date: str) -> int:
        existing = self._tables(conn)
        deleted = 0
        for spec in self.PARENT_SPECS:
            parent_rows = table_rows.get(spec["table"], [])
            parent_ids = [row["id"] for row in parent_rows if row.get("id") is not None]
            for child_table, fk in spec["children"]:
                if child_table in existing and parent_ids:
                    deleted += self._delete_by_ids(conn, child_table, fk, parent_ids)
        for spec in self.PARENT_SPECS:
            table = spec["table"]
            if table in existing:
                deleted += self._delete_between(conn, table, spec["date"], start_date, end_date)
        for table, date_column in self.DATE_SPECS:
            if table in existing:
                deleted += self._delete_between(conn, table, date_column, start_date, end_date)
        return deleted

    def _export_excel(self, table_rows: dict[str, list[dict[str, Any]]], path: Path, from_year: int, to_year: int) -> None:
        workbook = Workbook()
        info = workbook.active
        info.title = "Summary"
        info.append(["PRM Billing Inventory Free Space Export"])
        info.append(["Year From", from_year])
        info.append(["Year To", to_year])
        info.append(["Exported At", datetime.now().isoformat(timespec="seconds")])
        info.append([])
        info.append(["Table", "Rows"])
        used_titles = {"Summary"}
        for table, rows in sorted(table_rows.items()):
            info.append([table, len(rows)])
            sheet = workbook.create_sheet(self._sheet_title(table, used_titles))
            if not rows:
                sheet.append(["No rows"])
                continue
            headers = list(rows[0].keys())
            sheet.append(headers)
            for row in rows:
                sheet.append([row.get(header) for header in headers])
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)

    def _rows_between(self, conn: sqlite3.Connection, table: str, date_column: str, start_date: str, end_date: str) -> list[dict[str, Any]]:
        if not self._has_column(conn, table, date_column):
            return []
        rows = conn.execute(
            f"""
            SELECT *
            FROM "{table}"
            WHERE SUBSTR(COALESCE("{date_column}", ''), 1, 10) BETWEEN ? AND ?
            """,
            (start_date, end_date),
        ).fetchall()
        return [dict(row) for row in rows]

    def _rows_by_ids(self, conn: sqlite3.Connection, table: str, fk: str, ids: list[Any]) -> list[dict[str, Any]]:
        if not ids or not self._has_column(conn, table, fk):
            return []
        placeholders = ",".join("?" for _ in ids)
        rows = conn.execute(f'SELECT * FROM "{table}" WHERE "{fk}" IN ({placeholders})', tuple(ids)).fetchall()
        return [dict(row) for row in rows]

    def _delete_between(self, conn: sqlite3.Connection, table: str, date_column: str, start_date: str, end_date: str) -> int:
        if not self._has_column(conn, table, date_column):
            return 0
        cursor = conn.execute(
            f"""
            DELETE FROM "{table}"
            WHERE SUBSTR(COALESCE("{date_column}", ''), 1, 10) BETWEEN ? AND ?
            """,
            (start_date, end_date),
        )
        return int(cursor.rowcount or 0)

    def _delete_by_ids(self, conn: sqlite3.Connection, table: str, fk: str, ids: list[Any]) -> int:
        if not ids or not self._has_column(conn, table, fk):
            return 0
        placeholders = ",".join("?" for _ in ids)
        cursor = conn.execute(f'DELETE FROM "{table}" WHERE "{fk}" IN ({placeholders})', tuple(ids))
        return int(cursor.rowcount or 0)

    def _log(
        self,
        conn: sqlite3.Connection,
        from_year: int,
        to_year: int,
        backup_path: Path,
        export_path: Path,
        records_exported: int,
        records_deleted: int,
        user_name: str,
        table_rows: dict[str, list[dict[str, Any]]],
    ) -> None:
        if "data_maintenance_logs" not in self._tables(conn):
            return
        details = {
            "from_year": from_year,
            "to_year": to_year,
            "tables": {table: len(rows) for table, rows in table_rows.items() if rows},
        }
        conn.execute(
            """
            INSERT INTO data_maintenance_logs(
                operation,status,backup_path,archive_path,report_path,
                records_exported,records_deleted,created_by,details_json,created_at
            ) VALUES(?,?,?,?,?,?,?,?,?,?)
            """,
            (
                "free_space_year_range",
                "Success",
                str(backup_path),
                str(self.archive_root),
                str(export_path),
                records_exported,
                records_deleted,
                user_name,
                json.dumps(details, ensure_ascii=True),
                datetime.now().isoformat(timespec="seconds"),
            ),
        )

    def _tables(self, conn: sqlite3.Connection) -> set[str]:
        rows = conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'").fetchall()
        return {str(row[0]) for row in rows}

    def _has_column(self, conn: sqlite3.Connection, table: str, column: str) -> bool:
        return any(row[1] == column for row in conn.execute(f'PRAGMA table_info("{table}")'))

    def _sheet_title(self, table: str, used: set[str]) -> str:
        invalid_chars = set('[]:*?/\\')
        clean = "".join(ch if ch not in invalid_chars else "_" for ch in table)[:31] or "Table"
        title = clean
        suffix = 1
        while title in used:
            tail = f"_{suffix}"
            title = clean[: 31 - len(tail)] + tail
            suffix += 1
        used.add(title)
        return title
